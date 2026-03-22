# tests/test_ingest_excel.py
from datetime import date, datetime
from pathlib import Path
from openpyxl import Workbook
from post_trip_summary.pipeline.ingest.excel import ingest_excel


def _create_test_workbook(path: Path) -> None:
    wb = Workbook()
    ws_acc = wb.active
    ws_acc.title = "Accommodations"
    ws_acc.append(["Name", "Address", "City", "Country", "Check-In", "Check-Out"])
    ws_acc.append(["Hotel Le Marais", "123 Rue de Rivoli", "Paris", "France", "2026-03-05", "2026-03-08"])
    ws_transit = wb.create_sheet("Transit")
    ws_transit.append(["Mode", "From", "To", "Departure", "Arrival", "Details"])
    ws_transit.append(["flight", "JFK", "CDG", "2026-03-05 08:00", "2026-03-05 14:00", "AF001"])
    ws_act = wb.create_sheet("Activities")
    ws_act.append(["Name", "Location", "City", "Country", "Date", "Time", "Notes"])
    ws_act.append(["Seine River Cruise", "Port de la Bourdonnais", "Paris", "France", "2026-03-06", "19:00", "Reservation #12345"])
    ws_exp = wb.create_sheet("Expenses")
    ws_exp.append(["Date", "Amount", "Currency", "Merchant", "Category"])
    ws_exp.append(["2026-03-05", 87.50, "EUR", "REST LE PETIT", "dining"])
    wb.save(path)


def test_ingest_accommodations(tmp_path):
    wb_path = tmp_path / "trip.xlsx"
    _create_test_workbook(wb_path)
    result = ingest_excel(wb_path)
    assert len(result["accommodations"]) == 1
    acc = result["accommodations"][0]
    assert acc.name == "Hotel Le Marais"
    assert acc.check_in == datetime(2026, 3, 5, 0, 0)


def test_ingest_transit(tmp_path):
    wb_path = tmp_path / "trip.xlsx"
    _create_test_workbook(wb_path)
    result = ingest_excel(wb_path)
    assert len(result["transits"]) == 1
    t = result["transits"][0]
    assert t.mode == "flight"
    assert t.departure.name == "JFK"


def test_ingest_activities(tmp_path):
    wb_path = tmp_path / "trip.xlsx"
    _create_test_workbook(wb_path)
    result = ingest_excel(wb_path)
    assert len(result["activities"]) == 1
    act = result["activities"][0]
    assert act["name"] == "Seine River Cruise"


def test_ingest_expenses(tmp_path):
    wb_path = tmp_path / "trip.xlsx"
    _create_test_workbook(wb_path)
    result = ingest_excel(wb_path)
    assert len(result["expenses"]) == 1
    exp = result["expenses"][0]
    assert exp.amount == 87.50
    assert exp.currency == "EUR"


def test_ingest_missing_tab(tmp_path):
    wb_path = tmp_path / "minimal.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Accommodations"
    ws.append(["Name", "Address", "City", "Country", "Check-In", "Check-Out"])
    wb.save(wb_path)
    result = ingest_excel(wb_path)
    assert result["transits"] == []
    assert result["expenses"] == []


# --- New tests for plan Part 2 ---


def test_sheet_alias_matching(tmp_path):
    """Activities_Structured is found via alias for 'activities'."""
    wb_path = tmp_path / "alias.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Activities_Structured"
    ws.append(["Date", "Name", "City", "Country", "Type", "Time", "Notes"])
    ws.append(["2026-02-22", "Geothermal Park", "Rotorua", "New Zealand", "landmark", "", ""])
    wb.save(wb_path)
    result = ingest_excel(wb_path)
    assert len(result["activities"]) == 1
    assert result["activities"][0]["name"] == "Geothermal Park"


def test_activity_type_passthrough(tmp_path):
    """Activity type column is passed through when present."""
    wb_path = tmp_path / "typed.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Activities"
    ws.append(["Date", "Name", "City", "Country", "Type", "Time", "Notes"])
    ws.append(["2026-02-23", "Hobbiton", "Matamata", "New Zealand", "landmark", "09:30", ""])
    ws.append(["2026-02-26", "Roys Peak", "Wanaka", "New Zealand", "activity", "07:00", "6 hr hike"])
    wb.save(wb_path)
    result = ingest_excel(wb_path)
    assert len(result["activities"]) == 2
    assert result["activities"][0]["type"] == "landmark"
    assert result["activities"][1]["type"] == "activity"


def test_activity_without_type(tmp_path):
    """Activities without a type column don't have a type key."""
    wb_path = tmp_path / "notype.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Activities"
    ws.append(["Date", "Name", "City", "Country", "Time", "Notes"])
    ws.append(["2026-03-06", "Seine River Cruise", "Paris", "France", "19:00", ""])
    wb.save(wb_path)
    result = ingest_excel(wb_path)
    assert len(result["activities"]) == 1
    assert "type" not in result["activities"][0]


def test_credit_card_charges_parsing(tmp_path):
    """Credit Card Charges sheet is parsed as expenses with flag-column categories."""
    wb_path = tmp_path / "cc.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Credit Card Charges"
    ws.append(["Transaction Date", "Posted Date", "Card No.", "Description", "Category",
               "Debit", "Food", "Shopping", "Gas", "Parking", "Transportation", "Activities", "Lodging"])
    ws.append(["2026-03-04", "2026-03-05", 80, "Fergburger", "Dining", 21.76, "x", None, None, None, None, None, None])
    ws.append(["2026-03-05", "2026-03-06", 80, "NPD FRANKTON", "Gas/Automotive", 45.77, None, None, "x", None, None, None, None])
    ws.append(["2026-03-04", "2026-03-05", 80, "SKYLINE CABLEWAY", "Entertainment", 106.07, None, None, None, None, None, "x", None])
    wb.save(wb_path)
    result = ingest_excel(wb_path)
    assert len(result["expenses"]) == 3
    # Flag columns take priority
    assert result["expenses"][0].merchant == "Fergburger"
    assert result["expenses"][0].category == "dining"
    assert result["expenses"][0].amount == 21.76
    assert result["expenses"][0].currency == "NZD"
    assert result["expenses"][1].category == "gas"
    assert result["expenses"][2].category == "activity"


def test_credit_card_fallback_to_category_column(tmp_path):
    """When no flag columns are marked, falls back to Category column."""
    wb_path = tmp_path / "cc_fallback.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Credit Card Charges"
    ws.append(["Transaction Date", "Posted Date", "Card No.", "Description", "Category",
               "Debit", "Food", "Shopping", "Gas", "Parking", "Transportation", "Activities", "Lodging"])
    ws.append(["2026-03-01", "2026-03-02", 80, "SOME VENDOR", "Other Travel", 50.00,
               None, None, None, None, None, None, None])
    wb.save(wb_path)
    result = ingest_excel(wb_path)
    assert len(result["expenses"]) == 1
    assert result["expenses"][0].category == "other travel"


def test_expenses_budget_summary_skipped(tmp_path):
    """An expenses sheet that doesn't have date/amount columns is skipped gracefully."""
    wb_path = tmp_path / "budget.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"
    ws.append(["Expenses", None, None])
    ws.append([None, None, None])
    ws.append(["Lodging", "Location", "Nights"])
    ws.append([None, "Auckland", 1])
    wb.save(wb_path)
    result = ingest_excel(wb_path)
    assert result["expenses"] == []
