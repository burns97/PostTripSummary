# src/post_trip_summary/pipeline/ingest/excel.py
"""Parse trip itinerary from Excel workbook."""
import logging
from datetime import date, datetime
from pathlib import Path

import openpyxl

from post_trip_summary.models import (
    Accommodation, Transit, TransitPoint, Location, Expense,
)

log = logging.getLogger(__name__)

# Alias mapping: canonical name → list of acceptable sheet names (lowercase)
SHEET_ALIASES = {
    "accommodations": ["accommodations", "lodging"],
    "transit": ["transit", "transportation", "flights"],
    "activities": ["activities_structured", "activities", "schedule", "itinerary"],
    "expenses": ["expenses", "credit card charges", "credit_card_charges", "transactions"],
}

# Flag columns in Credit Card Charges sheets that map to expense categories
_CC_FLAG_COLUMNS = {
    "food": "dining",
    "shopping": "shopping",
    "gas": "gas",
    "parking": "parking",
    "transportation": "transportation",
    "activities": "activity",
    "lodging": "lodging",
}


def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def _parse_datetime(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip())
        except ValueError:
            pass
        try:
            return datetime.strptime(value.strip(), "%Y-%m-%d %H:%M")
        except ValueError:
            return None
    return None


def _rows_as_dicts(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    headers = [str(h).strip().lower().replace(" ", "_").replace("-", "_") for h in rows[0]]
    result = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        result.append(dict(zip(headers, row)))
    return result


def _find_sheet(wb, canonical: str) -> str | None:
    """Find the first matching sheet name for a canonical category."""
    sheet_names_lower = {s.lower(): s for s in wb.sheetnames}
    for alias in SHEET_ALIASES.get(canonical, [canonical]):
        if alias in sheet_names_lower:
            return sheet_names_lower[alias]
    return None


def _parse_accommodations(ws) -> list[Accommodation]:
    accommodations = []
    for row in _rows_as_dicts(ws):
        # Try datetime first (preserves time), fall back to date-only at midnight
        check_in = _parse_datetime(row.get("check_in"))
        if not check_in:
            d = _parse_date(row.get("check_in"))
            check_in = datetime.combine(d, datetime.min.time()) if d else None
        check_out = _parse_datetime(row.get("check_out"))
        if not check_out:
            d = _parse_date(row.get("check_out"))
            check_out = datetime.combine(d, datetime.min.time()) if d else None
        if not check_in or not check_out:
            continue
        loc = Location(lat=0.0, lon=0.0, name=str(row.get("name", "")), address=row.get("address"), city=str(row.get("city", "")), country=str(row.get("country", "")))
        accommodations.append(Accommodation(name=str(row.get("name", "")), location=loc, check_in=check_in, check_out=check_out, sources=["itinerary"]))
    return accommodations


def _parse_transits(ws) -> list[Transit]:
    transits = []
    for row in _rows_as_dicts(ws):
        dep_time = _parse_datetime(row.get("departure"))
        arr_time = _parse_datetime(row.get("arrival"))
        if not dep_time or not arr_time:
            continue
        dep_loc = Location(lat=0.0, lon=0.0, name=str(row.get("from", "")), address=None, city="", country="")
        arr_loc = Location(lat=0.0, lon=0.0, name=str(row.get("to", "")), address=None, city="", country="")
        transits.append(Transit(mode=str(row.get("mode", "unknown")).lower(), departure=TransitPoint(location=dep_loc, time=dep_time, name=str(row.get("from", ""))), arrival=TransitPoint(location=arr_loc, time=arr_time, name=str(row.get("to", ""))), details={"info": row.get("details", "")}, sources=["itinerary"]))
    return transits


def _parse_activities(ws) -> list[dict]:
    activities = []
    for row in _rows_as_dicts(ws):
        act_date = _parse_date(row.get("date"))
        if not act_date:
            continue
        activity = {
            "name": str(row.get("name", "")),
            "location": str(row.get("location", "")),
            "city": str(row.get("city", "")),
            "country": str(row.get("country", "")),
            "date": act_date.isoformat(),
            "time": str(row.get("time", "")),
            "notes": str(row.get("notes", "")),
        }
        # Pass through type if present
        raw_type = row.get("type")
        if raw_type:
            activity["type"] = str(raw_type).strip().lower()
        activities.append(activity)
    return activities


def _parse_expenses(ws) -> list[Expense]:
    """Parse a standard expenses sheet with date/amount/currency/merchant columns."""
    rows = _rows_as_dicts(ws)
    if not rows:
        return []
    # Check that the sheet actually has the expected columns
    first = rows[0]
    if "date" not in first or "amount" not in first:
        log.info("Expenses sheet does not have expected columns, skipping")
        return []
    expenses = []
    for row in rows:
        exp_date = _parse_date(row.get("date"))
        if not exp_date:
            continue
        expenses.append(Expense(date=exp_date, amount=float(row.get("amount", 0)), currency=str(row.get("currency", "USD")), merchant=str(row.get("merchant", "")), category=row.get("category"), source="spreadsheet"))
    return expenses


def _parse_credit_card_sheet(ws) -> list[Expense]:
    """Parse a Credit Card Charges sheet with flag columns for categories."""
    rows = _rows_as_dicts(ws)
    expenses = []
    for row in rows:
        exp_date = _parse_date(row.get("transaction_date"))
        if not exp_date:
            continue
        amount = row.get("debit")
        if amount is None:
            continue
        try:
            amount = float(amount)
        except (ValueError, TypeError):
            continue

        merchant = str(row.get("description", "")).strip()

        # Determine category from flag columns
        category = None
        for flag_col, cat_name in _CC_FLAG_COLUMNS.items():
            val = row.get(flag_col)
            if val and str(val).strip().lower() == "x":
                category = cat_name
                break
        # Fallback to Category column
        if not category:
            raw_cat = row.get("category")
            if raw_cat:
                category = str(raw_cat).strip().lower()

        expenses.append(Expense(
            date=exp_date,
            amount=amount,
            currency="NZD",
            merchant=merchant,
            category=category,
            source="credit_card",
        ))
    return expenses


def _find_and_parse_expenses(wb) -> list[Expense]:
    """Try each expense alias; return the first one that yields results."""
    sheet_names_lower = {s.lower(): s for s in wb.sheetnames}
    cc_names = {"credit card charges", "credit_card_charges"}
    for alias in SHEET_ALIASES["expenses"]:
        if alias not in sheet_names_lower:
            continue
        actual = sheet_names_lower[alias]
        if alias in cc_names:
            expenses = _parse_credit_card_sheet(wb[actual])
        else:
            expenses = _parse_expenses(wb[actual])
        if expenses:
            return expenses
    return []


def ingest_excel(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = {"accommodations": [], "transits": [], "activities": [], "expenses": []}

    acc_sheet = _find_sheet(wb, "accommodations")
    if acc_sheet:
        result["accommodations"] = _parse_accommodations(wb[acc_sheet])

    transit_sheet = _find_sheet(wb, "transit")
    if transit_sheet:
        result["transits"] = _parse_transits(wb[transit_sheet])

    act_sheet = _find_sheet(wb, "activities")
    if act_sheet:
        result["activities"] = _parse_activities(wb[act_sheet])

    # Try all expense sheet aliases; stop at the first one that yields results
    result["expenses"] = _find_and_parse_expenses(wb)

    wb.close()
    return result
